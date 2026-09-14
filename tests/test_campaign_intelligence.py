"""Comprehensive tests for AL AMR Step 14: Campaign Intelligence & Multi-Document Ingestion."""

from __future__ import annotations

import io
import json
from collections.abc import Iterator
from unittest.mock import patch, MagicMock
from pathlib import Path

import docx
import httpx
import pytest
from fastapi.testclient import TestClient

from autoclip import app as app_module
from autoclip.app import create_app
from autoclip.campaign import (
    CampaignConflict,
    CampaignNormalizer,
    CampaignSpecification,
    IngestedDocument,
    RequirementItem,
    extract_campaign_url,
)
from autoclip.campaign.conflict_detector import (
    check_aspect_ratio_conflict,
    check_banned_vs_required_conflict,
    check_cta_conflict,
    check_duration_conflict,
)
from autoclip.campaign.url_extractor import ExtractedUrlContent
from autoclip.db import store
from autoclip.db.models import CampaignSpecificationRecord, Job, Source, new_id
from autoclip.pipeline.ffmpeg import MediaInfo


@pytest.fixture
def client(autoclip_home, monkeypatch: pytest.MonkeyPatch) -> Iterator[TestClient]:
    monkeypatch.setenv(app_module.ENV_NO_WORKER, "1")
    with TestClient(create_app()) as test_client:
        yield test_client


def _make_docx(
    title: str = "Campaign Brief",
    text_paras: list[str] | None = None,
) -> bytes:
    doc = docx.Document()
    doc.add_heading(title, level=0)
    paras = text_paras or [
        "Must keep clips strictly between 30 and 60 seconds.",
        "Target audience: AI enthusiasts and tech founders.",
        "Mandatory requirement: Include strong Call to Action to subscribe at the end.",
        "Strictly prohibited: Do not discuss politics or cryptocurrency.",
    ]
    for p in paras:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf(text: str = "Campaign Rules\nDuration: 20-45s\nMust include brand logo") -> bytes:
    stream_content = f"BT\n/F1 12 Tf\n72 712 Td\n({text}) Tj\nET\n".encode("latin1")
    length = len(stream_content)
    return b"".join([
        b"%PDF-1.4\n",
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n",
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n",
        b"3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R/Resources<<\n/Font<</F1 4 0 R>>>>/Contents 5 0 R>>endobj\n",
        b"4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n",
        f"5 0 obj<</Length {length}>>stream\n".encode("ascii"),
        stream_content,
        b"endstream\nendobj\n",
        b"xref\n0 6\n0000000000 65535 f \n0000000009 00000 n \n0000000052 00000 n \n",
        b"0000000108 00000 n \n0000000213 00000 n \n0000000279 00000 n \n",
        b"trailer<</Size 6/Root 1 0 R>>\nstartxref\n443\n%%EOF",
    ])


# --------------------------------------------------------------------------
# 1. Models & RequirementItem Intelligence Tests
# --------------------------------------------------------------------------


def test_requirement_item_explicit_vs_inferred():
    explicit_item = RequirementItem(
        value=45.0,
        confidence="explicit",
        source_doc_id="doc_1",
        source_filename="brief_v2.pdf",
        snippet="Clips must strictly be between 30 and 60 seconds",
        weight=1.5,
    )
    assert explicit_item.confidence == "explicit"
    assert explicit_item.weight == 1.5

    d = explicit_item.to_dict()
    assert d["confidence"] == "explicit"
    assert d["value"] == 45.0

    restored = RequirementItem.from_dict(d)
    assert restored.value == 45.0
    assert restored.snippet == explicit_item.snippet


def test_ingested_document_provenance_and_hash():
    content = b"Mock document bytes for hashing"
    doc = IngestedDocument(
        filename="guideline.docx",
        source_type="docx",
        size_bytes=len(content),
        raw_text="Sample extracted text",
        word_count=3,
        char_count=21,
    )
    d = doc.to_dict()
    assert d["filename"] == "guideline.docx"
    assert d["status"] == "extracted"

    restored = IngestedDocument.from_dict(d)
    assert restored.doc_id == doc.doc_id
    assert restored.char_count == 21


# --------------------------------------------------------------------------
# 2. Conflict Detector Unit Tests
# --------------------------------------------------------------------------


def test_conflict_detection_duration_disjoint():
    doc_a = IngestedDocument(filename="client_deck.pdf")
    doc_b = IngestedDocument(filename="agency_brief.docx")

    # Disjoint ranges: [15, 30] vs [60, 90]
    conflict = check_duration_conflict(doc_a, (15.0, 30.0), doc_b, (60.0, 90.0))
    assert conflict is not None
    assert conflict.rule_category == "duration"
    assert conflict.severity == "critical"
    assert "Contradictory clip duration ranges" in conflict.description

    # Overlapping ranges: [20, 60] vs [30, 90] -> Overlap is [30, 60], no critical conflict
    conflict_overlap = check_duration_conflict(doc_a, (20.0, 60.0), doc_b, (30.0, 90.0))
    assert conflict_overlap is None


def test_conflict_detection_aspect_ratio():
    doc_a = IngestedDocument(filename="shorts_rule.pdf")
    doc_b = IngestedDocument(filename="landscape_rule.docx")

    conflict = check_aspect_ratio_conflict(doc_a, "9:16", doc_b, "16:9")
    assert conflict is not None
    assert conflict.rule_category == "aspect_ratio"
    assert conflict.severity == "critical"

    same = check_aspect_ratio_conflict(doc_a, "9:16", doc_b, "9:16")
    assert same is None


def test_conflict_detection_cta_contradiction():
    doc_a = IngestedDocument(filename="marketing.docx")
    doc_b = IngestedDocument(filename="editorial.pdf")

    conflict = check_cta_conflict(doc_a, True, doc_b, False, "strictly no call to action allowed")
    assert conflict is not None
    assert conflict.rule_category == "cta"
    assert "requires a Call to Action" in conflict.description


def test_conflict_detection_topic_collision():
    doc_a = IngestedDocument(filename="focus.pdf")
    doc_b = IngestedDocument(filename="brand_safety.docx")

    conflicts = check_banned_vs_required_conflict(
        doc_a, ["crypto trading", "machine learning"],
        doc_b, ["crypto trading", "gambling"],
    )
    assert len(conflicts) == 1
    assert conflicts[0].rule_category == "topics"
    assert "crypto trading" in conflicts[0].description


def test_version_superseding_auto_resolution():
    doc_old = IngestedDocument(filename="brief_v1.docx")
    doc_new = IngestedDocument(filename="brief_v2_final.docx")

    conflict = check_duration_conflict(doc_old, (15.0, 30.0), doc_new, (60.0, 90.0))
    assert conflict is not None
    # Version superseding should mark this superseded by the newer doc
    assert conflict.resolution_status == "superseded"
    assert "supersedes" in (conflict.resolution_notes or "").lower()


# --------------------------------------------------------------------------
# 3. CampaignNormalizer & Multi-Document Merging Tests
# --------------------------------------------------------------------------


def test_normalizer_single_and_multi_document():
    normalizer = CampaignNormalizer()

    docx_bytes = _make_docx(
        title="Creator Program Q4",
        text_paras=[
            "Strictly mandatory: Each clip must be 30 to 60 seconds.",
            "Must feature topic: Neural Networks and Computer Vision.",
            "Do not mention cryptocurrency or financial advice.",
            "Call to action: Check link in bio.",
        ],
    )
    pdf_bytes = _make_pdf("Target Audience: Tech Creators\nAspect Ratio: 9:16\nHook required within 3 seconds.")

    docs = normalizer.ingest_files([
        ("creator_q4.docx", docx_bytes),
        ("format_guide.pdf", pdf_bytes),
    ])

    assert len(docs) == 2
    assert docs[0].status == "extracted"
    assert docs[1].status == "extracted"

    spec = normalizer.normalize(docs)
    assert spec.title == "Creator Program Q4"
    assert spec.duration_min_s.value == 30.0
    assert spec.duration_max_s.value == 60.0
    assert spec.duration_min_s.confidence == "explicit"
    assert spec.aspect_ratio.value == "9:16"
    assert len(spec.documents) == 2
    assert any("Neural Networks" in item.value for item in spec.desired_topics)
    assert any("cryptocurrency" in item.value.lower() for item in spec.banned_topics)


def test_normalizer_partial_document_failure_non_blocking():
    normalizer = CampaignNormalizer()

    valid_docx = _make_docx("Valid Guide", ["Duration: 20-40s"])
    corrupt_bytes = b"CORRUPT_HEADER_NOT_A_VALID_DOCUMENT"

    docs = normalizer.ingest_files([
        ("corrupt.pdf", corrupt_bytes),
        ("valid.docx", valid_docx),
    ])

    assert len(docs) == 2
    assert docs[0].status == "failed"
    assert docs[0].error is not None
    assert docs[1].status == "extracted"

    # Normalization should succeed with remaining valid documents and record a warning
    spec = normalizer.normalize(docs)
    assert len(spec.documents) == 2
    assert len(spec.warnings) > 0
    assert "corrupt.pdf" in spec.warnings[0]
    assert spec.duration_min_s.value == 20.0


def test_campaign_specification_bridge_to_brief():
    normalizer = CampaignNormalizer()
    docx_bytes = _make_docx(
        "Bridge Test",
        [
            "Must keep duration between 25 and 55 seconds.",
            "Target audience: Developers.",
            "Output count: 7 clips.",
        ],
    )
    docs = normalizer.ingest_files([("test.docx", docx_bytes)])
    spec = normalizer.normalize(docs)

    brief = spec.to_campaign_brief()
    assert brief.name == "Bridge Test"
    assert brief.minimum_duration == 25.0
    assert brief.maximum_duration == 55.0
    assert brief.target_audience != ""


# --------------------------------------------------------------------------
# 4. Campaign URL Extractor & SSRF Protection Tests
# --------------------------------------------------------------------------


def test_extract_campaign_url_ssrf_rejection():
    # Loopback IP
    res_local = extract_campaign_url("http://127.0.0.1:8080/secret")
    assert res_local.status == "failed"
    assert "disallowed" in res_local.error.lower() or "blocked" in res_local.error.lower() or "invalid" in res_local.error.lower()

    # Private IP range
    res_private = extract_campaign_url("http://192.168.1.100/admin")
    assert res_private.status == "failed"

    # AWS metadata IP
    res_meta = extract_campaign_url("http://169.254.169.254/latest/meta-data/")
    assert res_meta.status == "failed"


def test_extract_campaign_url_success():
    html_content = """
    <!DOCTYPE html>
    <html>
      <head><title>Summer Clipping Contest - Whop</title></head>
      <body>
        <nav>Home | Login</nav>
        <main>
          <h1>Summer Clipping Contest</h1>
          <p>Create high-energy 9:16 TikTok and Shorts videos.</p>
          <p>Strictly mandatory: Clip length must be between 30 and 60 seconds.</p>
          <p>All clips must include the link in bio CTA.</p>
        </main>
        <footer>Copyright 2026</footer>
      </body>
    </html>
    """.encode("utf-8")

    mock_resp = httpx.Response(
        status_code=200,
        content=html_content,
        request=httpx.Request("GET", "https://example.com/campaign/summer"),
    )

    with patch("httpx.Client.get", return_value=mock_resp):
        res = extract_campaign_url("https://example.com/campaign/summer")
        assert res.status == "extracted"
        assert "Summer Clipping Contest" in res.title
        assert "30 and 60 seconds" in res.raw_text
        assert "Copyright" not in res.raw_text  # footer stripped


# --------------------------------------------------------------------------
# 5. Database Schema & Persistence Tests
# --------------------------------------------------------------------------


def test_campaign_specification_db_persistence(initialised_db):
    spec_id = new_id()
    spec_data = {
        "campaign_id": spec_id,
        "title": "Persistent Spec Test",
        "duration_min_s": {"value": 20.0, "confidence": "explicit"},
        "duration_max_s": {"value": 50.0, "confidence": "explicit"},
        "documents": [],
        "conflicts": [],
    }
    rec = CampaignSpecificationRecord(
        id=spec_id,
        job_id=None,
        title="Persistent Spec Test",
        spec=spec_data,
        has_conflicts=False,
        conflict_count=0,
        document_count=1,
    )
    store.create_campaign_spec(rec)

    fetched = store.get_campaign_spec(spec_id)
    assert fetched is not None
    assert fetched.title == "Persistent Spec Test"
    assert fetched.spec["duration_min_s"]["value"] == 20.0

    # Test update job_id
    mock_job_id = new_id()
    mock_source = Source(id=new_id(), type="upload", path=r"C:\test.mp4", filename="test.mp4")
    store.create_source(mock_source)
    store.create_job(Job(id=mock_job_id, source_id=mock_source.id, provider="test"))
    updated = store.update_campaign_spec(spec_id, job_id=mock_job_id)
    assert updated is not None
    assert updated.job_id == mock_job_id

    # Retrieve by job
    by_job = store.get_campaign_spec_for_job(mock_job_id)
    assert by_job is not None
    assert by_job.id == spec_id


# --------------------------------------------------------------------------
# 6. API Endpoints Integration Tests
# --------------------------------------------------------------------------


def test_api_campaign_intelligence_extract(client: TestClient):
    docx_bytes = _make_docx(
        "API Extract Test",
        ["Must be 15 to 45 seconds.", "Aspect ratio: 9:16."],
    )

    resp = client.post(
        "/api/campaigns/intelligence/extract",
        files=[("files", ("test.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["title"] == "API Extract Test"
    assert data["duration_min_s"]["value"] == 15.0
    assert data["duration_max_s"]["value"] == 45.0
    assert len(data["documents"]) == 1
    assert data["documents"][0]["filename"] == "test.docx"


def test_api_campaign_intelligence_url(client: TestClient):
    # Test rejection of private URL
    resp = client.post(
        "/api/campaigns/intelligence/url",
        json={"url": "http://127.0.0.1:9000/internal"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "failed"
    assert data["error"] is not None


def test_api_create_autonomous_job_with_campaign_spec(client: TestClient, tmp_path):
    # Create sample media file
    media_file = tmp_path / "test_source.mp4"
    media_file.write_bytes(b"\x00\x00\x00\x1cftypisom\x00\x00\x02\x00isomiso2mp41\x00\x00\x00\x08free")

    docx_bytes = _make_docx("Autonomous Step 14 Job", ["Clips must strictly be 30 to 60 seconds."])

    fake_media_info = MediaInfo(
        path=media_file,
        duration_s=120.0,
        width=1920,
        height=1080,
        fps=30.0,
        has_audio=True,
        has_video=True,
        audio_codec="aac",
        video_codec="h264",
    )

    with patch("autoclip.pipeline.ingest._probe_and_validate", return_value=fake_media_info):
        with open(media_file, "rb") as vf:
            resp = client.post(
                "/api/jobs/create-autonomous",
                files=[
                    ("video_file", ("test_source.mp4", vf, "video/mp4")),
                    ("guideline_files", ("campaign.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ],
                data={
                    "destinations": json.dumps(["telegram", "drive"]),
                },
            )

    assert resp.status_code == 201
    job_data = resp.json()
    assert job_data["id"] is not None
    assert job_data["campaign_spec_id"] is not None
    assert job_data["campaign_spec"] is not None
    assert job_data["campaign_spec"]["duration_min_s"]["value"] == 30.0

    # Retrieve campaign specification endpoint
    spec_resp = client.get(f"/api/jobs/{job_data['id']}/campaign-specification")
    assert spec_resp.status_code == 200
    spec_data = spec_resp.json()
    assert spec_data["title"] == "Autonomous Step 14 Job"
    assert len(spec_data["documents"]) >= 1
